import os
import tempfile
import unittest
from typing import Any

from pypdf import PdfReader, PdfWriter

from jot.services import document_service
from jot.services.document_service import DocumentError


def make_pdf(path: str, pages: int) -> str:
    writer = PdfWriter()
    for _ in range(pages):
        writer.add_blank_page(width=200, height=200)
    with open(path, "wb") as handle:
        writer.write(handle)
    return path


def make_pptx(path: str, slides: int, with_notes: bool = True) -> str:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    layout = prs.slide_layouts[1]
    for number in range(1, slides + 1):
        slide = prs.slides.add_slide(layout)
        title_shape: Any = slide.shapes.title
        title_shape.text = f"Titolo {number}"
        body_shape: Any = slide.placeholders[1]
        body = body_shape.text_frame
        body.text = f"Concetto principale {number}"
        para = body.add_paragraph()
        para.text = f"Dettaglio secondario {number}"
        para.level = 1
        if number == 2:
            rows, cols = 2, 2
            table = slide.shapes.add_table(
                rows, cols, Inches(1), Inches(4), Inches(4), Inches(1)
            ).table
            table.cell(0, 0).text = "A"
            table.cell(0, 1).text = "B"
            table.cell(1, 0).text = "1"
            table.cell(1, 1).text = "2"
        if with_notes:
            notes_frame: Any = slide.notes_slide.notes_text_frame
            notes_frame.text = f"Nota docente {number}"
    prs.save(path)
    return path


class ParsePageRangesTests(unittest.TestCase):
    def test_empty_selects_all(self):
        self.assertEqual(document_service.parse_page_ranges("", 4), [1, 2, 3, 4])

    def test_mixed_ranges_sorted_unique(self):
        self.assertEqual(
            document_service.parse_page_ranges("5-7, 2; 6 , 9\u201310", 12),
            [2, 5, 6, 7, 9, 10],
        )

    def test_reversed_range_is_normalized(self):
        self.assertEqual(document_service.parse_page_ranges("8-6", 10), [6, 7, 8])

    def test_out_of_bounds_rejected(self):
        with self.assertRaises(DocumentError):
            document_service.parse_page_ranges("1-50", 10)

    def test_garbage_rejected(self):
        with self.assertRaises(DocumentError):
            document_service.parse_page_ranges("capitolo 3", 10)

    def test_format_roundtrip(self):
        pages = document_service.parse_page_ranges("1-3,7,9-10", 20)
        self.assertEqual(document_service.format_page_ranges(pages), "1-3, 7, 9-10")


class PlanSlideChunksTests(unittest.TestCase):
    def test_even_split(self):
        self.assertEqual(
            document_service.plan_slide_chunks(24, 12), [(1, 12), (13, 24)]
        )

    def test_small_tail_merged(self):
        # 25 slides / 12 -> last block would contain 1 slide: merged
        self.assertEqual(
            document_service.plan_slide_chunks(25, 12), [(1, 12), (13, 25)]
        )

    def test_reasonable_tail_kept(self):
        self.assertEqual(
            document_service.plan_slide_chunks(30, 12),
            [(1, 12), (13, 24), (25, 30)],
        )

    def test_block_size_clamped(self):
        self.assertEqual(document_service.plan_slide_chunks(5, 1), [(1, 3), (4, 5)])
        self.assertEqual(document_service.plan_slide_chunks(0, 12), [])

    def test_single_block(self):
        self.assertEqual(document_service.plan_slide_chunks(7, 12), [(1, 7)])


class PdfTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)

    def test_validate_slides_counts_pages(self):
        path = make_pdf(os.path.join(self.tmp.name, "slide.pdf"), 9)
        self.assertEqual(document_service.validate_slides(path), 9)

    def test_unsupported_extension(self):
        path = os.path.join(self.tmp.name, "slide.ppt")
        with open(path, "wb") as handle:
            handle.write(b"old")
        with self.assertRaises(DocumentError) as ctx:
            document_service.validate_slides(path)
        self.assertIn(".pptx", str(ctx.exception))

    def test_corrupted_pdf(self):
        path = os.path.join(self.tmp.name, "broken.pdf")
        with open(path, "wb") as handle:
            handle.write(b"not a pdf at all")
        with self.assertRaises(DocumentError):
            document_service.validate_slides(path)

    def test_slide_chunk_contains_only_requested_pages(self):
        path = make_pdf(os.path.join(self.tmp.name, "slide.pdf"), 10)
        payload = document_service.build_slide_chunk(path, 4, 7)
        self.assertIsNotNone(payload.pdf_bytes)
        sub = os.path.join(self.tmp.name, "sub.pdf")
        with open(sub, "wb") as handle:
            handle.write(payload.pdf_bytes or b"")
        self.assertEqual(len(PdfReader(sub).pages), 4)
        self.assertEqual((payload.start, payload.end), (4, 7))

    def test_book_excerpt_and_limit(self):
        book = make_pdf(os.path.join(self.tmp.name, "libro.pdf"), 20)
        total, selected = document_service.validate_book(book, "3-5, 10")
        self.assertEqual(total, 20)
        out = os.path.join(self.tmp.name, "excerpt.pdf")
        self.assertEqual(document_service.write_book_excerpt(book, selected, out), 4)
        self.assertEqual(len(PdfReader(out).pages), 4)
        self.assertFalse(os.path.exists(out + ".partial"))

    def test_book_too_many_pages(self):
        book = make_pdf(
            os.path.join(self.tmp.name, "big.pdf"), document_service.MAX_BOOK_PAGES + 5
        )
        with self.assertRaises(DocumentError) as ctx:
            document_service.validate_book(book, "")
        self.assertIn("massimo", str(ctx.exception))
        total, selected = document_service.validate_book(book, "1-10")
        self.assertEqual(len(selected), 10)
        self.assertEqual(total, document_service.MAX_BOOK_PAGES + 5)

    def test_book_must_be_pdf(self):
        path = make_pptx(os.path.join(self.tmp.name, "libro.pptx"), 2)
        with self.assertRaises(DocumentError):
            document_service.validate_book(path, "")


class PptxTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.path = make_pptx(os.path.join(self.tmp.name, "lezione.pptx"), 5)

    def test_counts_slides(self):
        self.assertEqual(document_service.validate_slides(self.path), 5)

    def test_chunk_text_extraction(self):
        payload = document_service.build_slide_chunk(self.path, 2, 3)
        self.assertIsNone(payload.pdf_bytes)
        self.assertIn("### Slide 2: Titolo 2", payload.text)
        self.assertIn("### Slide 3: Titolo 3", payload.text)
        self.assertNotIn("Slide 1", payload.text)
        self.assertIn("- Concetto principale 2", payload.text)
        self.assertIn("  - Dettaglio secondario 2", payload.text)
        self.assertIn("Note del docente: Nota docente 3", payload.text)
        self.assertIn("| A | B |", payload.text)
        self.assertEqual(payload.images, [])


if __name__ == "__main__":
    unittest.main()
