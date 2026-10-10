"""
Document helpers for the "Slide e appunti" mode.

Reads lesson slides (PDF or PPTX) and an optional textbook (PDF), splits the
slides into blocks and prepares the payloads sent to Gemini.  Everything here is
pure local file processing: no network calls.
"""

from __future__ import annotations

import io
import os
import re
from dataclasses import dataclass, field

SLIDE_EXTS: frozenset[str] = frozenset({".pdf", ".pptx"})
BOOK_EXTS: frozenset[str] = frozenset({".pdf"})

DEFAULT_SLIDES_PER_CHUNK = 12
MIN_SLIDES_PER_CHUNK = 3
MAX_SLIDES_PER_CHUNK = 40

# Pages of the textbook sent along with every block.  Above this the request
# becomes too heavy for the free tier (tokens per minute) and too slow.
MAX_BOOK_PAGES = 300

# Gemini rejects inline requests above ~20 MB: keep a safety margin and switch
# to the Files API above this threshold.
INLINE_MAX_BYTES = 14 * 1024 * 1024

# PPTX picture handling
_PPTX_IMAGE_MIMES: frozenset[str] = frozenset(
    {"image/png", "image/jpeg", "image/jpg", "image/webp", "image/gif"}
)
_PPTX_MAX_IMAGES_PER_SLIDE = 3
_PPTX_MAX_IMAGES_PER_CHUNK = 16
_PPTX_MAX_IMAGE_BYTES = 4 * 1024 * 1024


class DocumentError(ValueError):
    """Raised when a slide/book file cannot be used. Message is user-facing (IT)."""


@dataclass(slots=True)
class SlideChunkPayload:
    """Content of one block of slides ready to be turned into Gemini parts."""

    start: int
    end: int
    pdf_bytes: bytes | None = None
    text: str = ""
    images: list[tuple[bytes, str]] = field(default_factory=list)

    @property
    def total_bytes(self) -> int:
        size = len(self.pdf_bytes or b"") + len(self.text.encode("utf-8"))
        return size + sum(len(data) for data, _mime in self.images)


def file_ext(path: str) -> str:
    return os.path.splitext(str(path or ""))[1].lower()


def is_slides_path(path: str) -> bool:
    return file_ext(path) in SLIDE_EXTS


def is_book_path(path: str) -> bool:
    return file_ext(path) in BOOK_EXTS


def clamp_slides_per_chunk(value) -> int:
    try:
        number = int(value)
    except (TypeError, ValueError):
        return DEFAULT_SLIDES_PER_CHUNK
    return max(MIN_SLIDES_PER_CHUNK, min(MAX_SLIDES_PER_CHUNK, number))


# ---------------------------------------------------------------------------
# Counting / validation
# ---------------------------------------------------------------------------


def _open_pdf(path: str):
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover - dependency missing
        raise DocumentError("Componente PDF non disponibile (pypdf).") from exc
    try:
        reader = PdfReader(path)
        if reader.is_encrypted:
            try:
                ok = reader.decrypt("")
            except Exception:
                ok = 0
            if not ok:
                raise DocumentError(
                    f"Il PDF '{os.path.basename(path)}' è protetto da password."
                )
        return reader
    except DocumentError:
        raise
    except Exception as exc:
        raise DocumentError(
            f"Impossibile leggere il PDF '{os.path.basename(path)}': file danneggiato o non valido."
        ) from exc


def _open_pptx(path: str):
    try:
        from pptx import Presentation
    except ImportError as exc:  # pragma: no cover - dependency missing
        raise DocumentError(
            "Componente PowerPoint non disponibile (python-pptx)."
        ) from exc
    try:
        return Presentation(path)
    except Exception as exc:
        raise DocumentError(
            f"Impossibile leggere la presentazione '{os.path.basename(path)}': "
            "file danneggiato o non in formato .pptx."
        ) from exc


def count_pages(path: str) -> int:
    """Return the number of slides (PPTX) or pages (PDF)."""
    ext = file_ext(path)
    if not os.path.isfile(path):
        raise DocumentError("File non trovato.")
    if ext == ".pdf":
        return len(_open_pdf(path).pages)
    if ext == ".pptx":
        return len(_open_pptx(path).slides)
    raise DocumentError("Formato non supportato. Usa un file PDF o PPTX.")


def validate_slides(path: str) -> int:
    if not is_slides_path(path):
        raise DocumentError(
            "Formato slide non supportato. Usa un PDF o una presentazione PowerPoint (.pptx). "
            "Se hai un .ppt, salvalo prima come .pptx o esportalo in PDF."
        )
    pages = count_pages(path)
    if pages <= 0:
        raise DocumentError("Il file delle slide non contiene pagine.")
    return pages


def validate_book(path: str, page_spec: str = "") -> tuple[int, list[int]]:
    """Validate the textbook and the page selection. Returns (total_pages, selected)."""
    if not is_book_path(path):
        raise DocumentError("Il libro deve essere un file PDF.")
    total = count_pages(path)
    if total <= 0:
        raise DocumentError("Il PDF del libro non contiene pagine.")
    selected = parse_page_ranges(page_spec, total)
    if len(selected) > MAX_BOOK_PAGES:
        raise DocumentError(
            f"Hai selezionato {len(selected)} pagine del libro: il massimo è {MAX_BOOK_PAGES}. "
            "Indica solo i capitoli della lezione (es. 120-165)."
        )
    return total, selected


def parse_page_ranges(spec: str, total_pages: int) -> list[int]:
    """Parse a selection like "12-30, 45, 50-52" into sorted 1-based page numbers.

    An empty spec selects every page.  Raises DocumentError on invalid input.
    """
    total = int(total_pages)
    text = str(spec or "").strip()
    if not text:
        return list(range(1, total + 1))
    pages: set[int] = set()
    for raw_part in re.split(r"[,;]", text):
        part = raw_part.strip()
        if not part:
            continue
        match = re.fullmatch(r"(\d+)\s*(?:-|\u2013|\u2014)\s*(\d+)", part)
        if match:
            first, last = int(match.group(1)), int(match.group(2))
        elif re.fullmatch(r"\d+", part):
            first = last = int(part)
        else:
            raise DocumentError(
                f"Selezione pagine non valida: '{part}'. Usa un formato come 12-30, 45."
            )
        if first > last:
            first, last = last, first
        if first < 1 or last > total:
            raise DocumentError(
                f"Pagine {first}-{last} fuori dal documento (il libro ha {total} pagine)."
            )
        pages.update(range(first, last + 1))
    if not pages:
        raise DocumentError("Nessuna pagina selezionata.")
    return sorted(pages)


def format_page_ranges(pages: list[int]) -> str:
    """Compact representation: [1,2,3,7] -> "1-3, 7"."""
    if not pages:
        return ""
    ordered = sorted(set(int(p) for p in pages))
    ranges: list[str] = []
    start = prev = ordered[0]
    for page in ordered[1:]:
        if page == prev + 1:
            prev = page
            continue
        ranges.append(f"{start}-{prev}" if start != prev else str(start))
        start = prev = page
    ranges.append(f"{start}-{prev}" if start != prev else str(start))
    return ", ".join(ranges)


def plan_slide_chunks(
    total_slides: int, slides_per_chunk: int
) -> list[tuple[int, int]]:
    """Split 1..total into inclusive (start, end) blocks.

    A tiny trailing block (less than a third of the block size) is merged into
    the previous one so the model never receives a single orphan slide.
    """
    total = int(total_slides)
    size = clamp_slides_per_chunk(slides_per_chunk)
    if total <= 0:
        return []
    chunks: list[tuple[int, int]] = []
    start = 1
    while start <= total:
        end = min(total, start + size - 1)
        chunks.append((start, end))
        start = end + 1
    if len(chunks) >= 2:
        last_start, last_end = chunks[-1]
        if (last_end - last_start + 1) < max(2, size // 3):
            prev_start, _prev_end = chunks[-2]
            chunks[-2:] = [(prev_start, last_end)]
    return chunks


# ---------------------------------------------------------------------------
# Payload builders
# ---------------------------------------------------------------------------


def build_pdf_subset(path: str, pages: list[int]) -> bytes:
    """Return a new PDF containing only the given 1-based pages."""
    from pypdf import PdfWriter

    reader = _open_pdf(path)
    writer = PdfWriter()
    total = len(reader.pages)
    for page_no in pages:
        if 1 <= int(page_no) <= total:
            writer.add_page(reader.pages[int(page_no) - 1])
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def write_book_excerpt(book_path: str, pages: list[int], out_path: str) -> int:
    """Write the selected book pages to out_path. Returns the number of pages."""
    data = build_pdf_subset(book_path, pages)
    tmp_path = f"{out_path}.partial"
    with open(tmp_path, "wb") as handle:
        handle.write(data)
    os.replace(tmp_path, out_path)
    return len(pages)


def _shape_text(shape) -> list[str]:
    lines: list[str] = []
    try:
        if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
            for paragraph in shape.text_frame.paragraphs:
                text = "".join(run.text for run in paragraph.runs).strip()
                if text:
                    indent = "  " * int(getattr(paragraph, "level", 0) or 0)
                    lines.append(f"{indent}- {text}")
        if getattr(shape, "has_table", False) and shape.has_table:
            for row in shape.table.rows:
                cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                lines.append("| " + " | ".join(cells) + " |")
        if getattr(shape, "shape_type", None) is not None and hasattr(shape, "shapes"):
            for child in shape.shapes:  # group shapes
                lines.extend(_shape_text(child))
    except Exception:
        pass
    return lines


def _shape_images(shape) -> list[tuple[bytes, str]]:
    images: list[tuple[bytes, str]] = []
    try:
        image = getattr(shape, "image", None)
        if image is not None:
            mime = str(getattr(image, "content_type", "") or "").lower()
            blob = image.blob
            if mime in _PPTX_IMAGE_MIMES and 0 < len(blob) <= _PPTX_MAX_IMAGE_BYTES:
                images.append((blob, "image/jpeg" if mime == "image/jpg" else mime))
        if hasattr(shape, "shapes"):
            for child in shape.shapes:
                images.extend(_shape_images(child))
    except Exception:
        pass
    return images


def build_pptx_chunk(path: str, start: int, end: int) -> SlideChunkPayload:
    """Extract text, speaker notes and pictures of slides start..end (1-based)."""
    presentation = _open_pptx(path)
    slides = list(presentation.slides)
    sections: list[str] = []
    images: list[tuple[bytes, str]] = []
    for number in range(int(start), int(end) + 1):
        if number < 1 or number > len(slides):
            continue
        slide = slides[number - 1]
        title = ""
        try:
            if slide.shapes.title is not None:
                title = slide.shapes.title.text.strip()
        except Exception:
            title = ""
        body: list[str] = []
        slide_images: list[tuple[bytes, str]] = []
        for shape in slide.shapes:
            try:
                if (
                    title
                    and slide.shapes.title is not None
                    and shape == slide.shapes.title
                ):
                    continue
            except Exception:
                pass
            body.extend(_shape_text(shape))
            slide_images.extend(_shape_images(shape))
        notes = ""
        try:
            if slide.has_notes_slide:
                frame = slide.notes_slide.notes_text_frame
                notes = frame.text.strip() if frame is not None else ""
        except Exception:
            notes = ""
        header = f"### Slide {number}" + (f": {title}" if title else "")
        parts = [header]
        parts.append("\n".join(body) if body else "(nessun testo nella slide)")
        if notes:
            parts.append(f"Note del docente: {notes}")
        kept = slide_images[:_PPTX_MAX_IMAGES_PER_SLIDE]
        room = _PPTX_MAX_IMAGES_PER_CHUNK - len(images)
        kept = kept[: max(0, room)]
        if kept:
            first_index = len(images) + 1
            last_index = len(images) + len(kept)
            label = (
                f"immagine {first_index}"
                if first_index == last_index
                else f"immagini {first_index}-{last_index}"
            )
            parts.append(f"(Contiene {label} allegate in ordine.)")
            images.extend(kept)
        sections.append("\n".join(parts))
    return SlideChunkPayload(
        start=int(start),
        end=int(end),
        text="\n\n".join(sections).strip(),
        images=images,
    )


def build_slide_chunk(path: str, start: int, end: int) -> SlideChunkPayload:
    ext = file_ext(path)
    if ext == ".pdf":
        data = build_pdf_subset(path, list(range(int(start), int(end) + 1)))
        return SlideChunkPayload(start=int(start), end=int(end), pdf_bytes=data)
    if ext == ".pptx":
        return build_pptx_chunk(path, start, end)
    raise DocumentError("Formato slide non supportato.")
